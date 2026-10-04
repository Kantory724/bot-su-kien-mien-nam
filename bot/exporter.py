"""Xuất danh sách sự kiện ra Excel theo tuần/tháng."""
import calendar
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from . import config
from .db import DB
from .geo import province_name
from .models import PRIORITY_LABEL

FILL = {"CAO": "F8CBAD", "TB": "FFE699", "THAP": "E2EFDA"}


def period_range(period: str, today: date) -> tuple[date, date]:
    if period == "month":
        return today.replace(day=1), today.replace(day=calendar.monthrange(today.year, today.month)[1])
    return today, today + timedelta(days=6)


def export_excel(db: DB, period: str = "week", out: str | None = None, today: date | None = None) -> Path:
    today = today or config.today()
    a, b = period_range(period, today)
    events = db.events_between(a, b)
    wb = Workbook()
    ws = wb.active
    ws.title = "Sự kiện"
    head = ["Từ ngày", "Đến ngày", "Giờ", "Tỉnh/TP", "Tên sự kiện", "Địa điểm", "Quy mô (người)",
            "Pháo hoa", "Đại nhạc hội", "Mức ưu tiên", "Tóm tắt", "Nguồn"]
    ws.append(head)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F4E78")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for e in events:
        ws.append([e.start_date, e.last_date, e.start_time, province_name(e.province), e.name, e.venue,
                   e.crowd, "Có" if e.fireworks else "", "Có" if e.big_concert else "",
                   PRIORITY_LABEL[e.priority], e.summary, e.sources[0] if e.sources else ""])
        row = ws.max_row
        ws.cell(row, 10).fill = PatternFill("solid", fgColor=FILL[e.priority])
        for col in (1, 2):
            ws.cell(row, col).number_format = "DD/MM/YYYY"
        if e.sources:
            ws.cell(row, 12).hyperlink = e.sources[0]
            ws.cell(row, 12).font = Font(color="0563C1", underline="single")
        ws.cell(row, 7).number_format = "#,##0"
    for i, w in enumerate([12, 12, 7, 16, 48, 36, 14, 9, 12, 11, 60, 50], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    out_path = Path(out) if out else config.ROOT / "out" / f"su-kien-{period}-{a:%Y%m%d}.xlsx"
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
