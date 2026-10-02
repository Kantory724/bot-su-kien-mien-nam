from datetime import date, datetime, timedelta

import pytest

from bot import config
from bot.collector import Item, load_sources
from bot.commands import handle
from bot.db import DB
from bot.exporter import export_excel
from bot.extractor import (extract_rules, find_crowd, find_dates, find_time, find_venue, has_fireworks,
                           event_keyword_hits)
from bot.formatter import fmt_range, format_digest
from bot.geo import detect_province, resolve_province
from bot.models import Event
from bot.pipeline import ingest, is_duplicate, notify_source_failures, run_digest
from bot.telegram import split_message

REF = date(2026, 10, 2)


def d(s):
    return date.fromisoformat(s)


# ---------- ngày/giờ/quy mô ----------
@pytest.mark.parametrize("text,exp", [
    ("từ ngày 17 đến 19/10 tại", ("2026-10-17", "2026-10-19")),
    ("ngày 15 tháng 10", ("2026-10-15", "2026-10-15")),
    ("17/10-18/10/2026", ("2026-10-17", "2026-10-18")),
    ("17-18/10", ("2026-10-17", "2026-10-18")),
    ("17 và 18/10", ("2026-10-17", "2026-10-18")),
    ("từ 28 tháng 10 đến ngày 2 tháng 11 năm 2026", ("2026-10-28", "2026-11-02")),
    ("khai mạc ngày 2/9", ("2026-09-02", "2026-09-02")),
])
def test_dates(text, exp):
    got = find_dates(text, REF)
    assert (got[0][1].isoformat(), got[0][2].isoformat()) == exp


def test_lunar_skipped_and_year_rollover():
    got = find_dates("ngày 15 tháng 10 âm lịch (tức 24/11)", REF)
    assert [g[1] for g in got] == [d("2026-11-24")]
    assert find_dates("ngày 1/1", date(2026, 12, 20))[0][1] == d("2027-01-01")
    g = find_dates("31/12 đến 1/1", date(2026, 12, 20))[0]
    assert (g[1], g[2]) == (d("2026-12-31"), d("2027-01-01"))
    assert find_dates("ngày 31/2", REF) == []  # ngày không tồn tại


def test_time():
    assert find_time("lúc 19h30") == "19:30"
    assert find_time("từ 20 giờ") == "20:00"
    assert find_time("rộng 10ha") == ""
    assert find_time("mở cửa 18:00") == "18:00"


@pytest.mark.parametrize("text,exp", [
    ("hơn 20.000 người", 20000), ("15 nghìn lượt khách", 15000), ("1,5 triệu lượt khách", 1500000),
    ("hàng chục nghìn người", 20000), ("đón 5 triệu lượt khách năm 2025", None), ("không có số liệu", None),
])
def test_crowd(text, exp):
    assert find_crowd(text) == exp


def test_flags_venue():
    assert has_fireworks("sẽ bắn pháo hoa") and not has_fireworks("không bắn pháo hoa năm nay")
    assert find_venue("diễn ra tại Công viên Văn hóa Suối Tiên vào lúc 19h") == "Công viên Văn hóa Suối Tiên"
    assert find_venue("tổ chức tại Cần Thơ") == ""  # tên tỉnh không phải địa điểm cụ thể
    assert "lễ hội" in event_keyword_hits("Khai mạc Lễ hội Ok Om Bok") or "le hoi" in event_keyword_hits("Lễ hội Ok Om Bok")


# ---------- tỉnh ----------
def test_province():
    assert detect_province("Chủ tịch Hồ Chí Minh", "Hồ Chí Minh nói") is None
    assert detect_province("Đại nhạc hội tại TP.HCM") == "hcm"
    assert detect_province("Lễ hội ở Phú Quốc") == "angiang"          # Kiên Giang đã sáp nhập An Giang
    assert detect_province("Hội chợ Bến Tre") == "vinhlong"
    assert detect_province("Tri ân thầy cô") is None                  # 'tri ân' không bị nhận nhầm
    assert resolve_province("Cần Thơ") == "cantho" and resolve_province("hcm") == "hcm"
    assert resolve_province("xyz") is None


# ---------- khử trùng ----------
def test_dedupe():
    a = Event("Đại nhạc hội Mùa Hè Xanh tại Cần Thơ thu hút 20.000 khán giả", "cantho", start_date=d("2026-10-07"))
    b = Event("Cần Thơ: Mùa Hè Xanh - đại nhạc hội lớn nhất năm", "cantho", start_date=d("2026-10-07"))
    c = Event("Lễ hội Vía Bà Chúa Xứ", "cantho", start_date=d("2026-10-07"))
    e = Event("Đại nhạc hội Mùa Hè Xanh", "camau", start_date=d("2026-10-07"))
    assert is_duplicate(a, b) and not is_duplicate(a, c) and not is_duplicate(a, e)


# ---------- pipeline ----------
def make_item(title, link="https://x/1", hint="cantho"):
    return Item("t", "T", title, link, "", REF, hint)


def test_ingest_flow():
    db = DB(":memory:")
    body = "Đại nhạc hội diễn ra ngày 17-18/10 tại Quảng trường Hòa Bình, phường Ninh Kiều, dự kiến 20.000 người, bắn pháo hoa lúc 21h."
    assert ingest(db, make_item("Đại nhạc hội Cần Thơ", "u1"), body, None, REF) == "new"
    assert ingest(db, make_item("Cần Thơ có đại nhạc hội lớn", "u2"), body, None, REF) in ("merged", "dup")
    assert db.count_events() == 1
    ev = db.events_between(d("2026-10-17"), d("2026-10-17"))[0]
    assert ev.priority == "CAO" and ev.is_large and ev.crowd == 20000 and ev.fireworks
    assert ingest(db, make_item("Lễ hội xưa", "u3", "dongnai"), "Lễ hội ngày 10/9.", None, REF) == "past"
    assert ingest(db, make_item("Lễ hội Hà Nội", "u4", None), "tại Hà Nội ngày 20/10", None, REF) == "out_of_scope"
    assert [e.id for e in db.unalerted_large(REF)] == [ev.id]
    db.mark_alerted(ev.id)
    assert db.unalerted_large(REF) == []


def test_priority_levels():
    assert Event("Lễ hội X", "hcm", crowd=12000).compute_priority() == "CAO"
    assert Event("Hội chợ Y", "hcm").compute_priority() == "TB"
    assert Event("Giao lưu nhỏ", "hcm").compute_priority() == "THAP"
    assert Event("Đêm nhạc Z", "hcm", fireworks=True).compute_priority() == "CAO"


# ---------- định dạng ----------
def test_format():
    assert fmt_range(d("2026-10-17"), d("2026-10-18")) == "17–18/10"
    assert fmt_range(d("2026-10-28"), d("2026-11-02")) == "28/10–2/11"
    assert fmt_range(d("2026-10-17"), d("2026-10-17")) == "17/10"
    txt = format_digest([], d("2026-10-12"), 7)
    assert "Thứ Hai 12/10/2026" in txt
    long = "\n\n".join("x" * 500 for _ in range(30))
    parts = split_message(long)
    assert all(len(p) <= 3900 for p in parts) and len(parts) > 1


# ---------- lệnh ----------
def seeded_db():
    db = DB(":memory:")
    t = config.today()
    db.insert_event(Event("Lễ hội A", "cantho", "Bến Ninh Kiều", t, t + timedelta(days=1), crowd=20000, sources=["http://a"]))
    db.insert_event(Event("Hội chợ B pháo hoa", "tayninh", "TT Hội chợ", t + timedelta(days=3), t + timedelta(days=3), sources=["http://b"]))
    db.insert_event(Event("Giải chạy C", "camau", "", t + timedelta(days=30), t + timedelta(days=30), sources=["http://c"]))
    return db


def test_commands():
    db = seeded_db()
    assert "Lễ hội A" in handle("/homnay", "1", db).texts[0]
    w = handle("/tuannay", "1", db).texts[0]
    assert "Lễ hội A" in w and "Hội chợ B" in w and "Giải chạy C" not in w
    assert "Lễ hội A" in handle("/tinh Cần Thơ", "1", db).texts[0]
    assert "Giải chạy C" in handle("/tinh ca mau", "1", db).texts[0]
    assert "Không nhận ra" in handle("/tinh Hà Nội", "1", db).texts[0]
    assert "Hội chợ B" in handle("/sukien phao hoa", "1", db).texts[0]
    assert "Hội chợ B" in handle("/sukien@my_bot pháo hoa", "1", db).texts[0]
    assert handle("/excel thang", "1", db).excel == "month"
    assert "Chat ID" in handle("/id", "42", db).texts[0]
    assert "không hợp lệ" in handle("/abc", "1", db).texts[0]


# ---------- Excel ----------
def test_excel(tmp_path):
    from openpyxl import load_workbook
    p = export_excel(seeded_db(), "week", str(tmp_path / "a.xlsx"))
    ws = load_workbook(p).active
    assert ws.max_row == 3 and ws["E2"].value == "Lễ hội A" and ws["J2"].value == "CAO"


# ---------- Telegram giả: poll, digest, cảnh báo nguồn lỗi, lịch ----------
class FakeTG:
    def __init__(self, updates=None):
        self.sent, self.docs, self.updates = [], [], updates or []
        self.dry = False

    def get_updates(self, offset, timeout=0):
        return [u for u in self.updates if not offset or u["update_id"] >= offset]

    def send(self, chat_id, text):
        self.sent.append((chat_id, text))

    def broadcast(self, text, chat_ids=None):
        ids = chat_ids if chat_ids is not None else config.chat_ids()
        for c in ids:
            self.send(c, text)
        return len(ids)

    def send_document(self, chat_id, path, caption=""):
        self.docs.append((chat_id, str(path)))


def upd(i, chat, text):
    return {"update_id": i, "message": {"text": text, "chat": {"id": chat}}}


def test_poll_authorization(monkeypatch):
    from bot.runner import poll_once
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "100")
    db = seeded_db()
    tg = FakeTG([upd(1, 100, "/homnay"), upd(2, 999, "/homnay"), upd(3, 999, "/id"), upd(4, 100, "/excel")])
    poll_once(db, tg)
    out = {(c, t[:25]) for c, t in tg.sent}
    assert any(c == "100" and "SỰ KIỆN HÔM NAY" in t for c, t in tg.sent)
    assert any(c == "999" and "chưa được cấp quyền" in t for c, t in tg.sent)
    assert any(c == "999" and "Chat ID" in t for c, t in tg.sent)
    assert len(tg.docs) == 1 and tg.docs[0][0] == "100"
    assert db.get("tg_offset") == "5"
    n = len(tg.sent)
    poll_once(db, tg)  # không xử lý lại tin cũ
    assert len(tg.sent) == n


def test_digest_once_per_day(monkeypatch):
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "100,200")
    db, tg = seeded_db(), FakeTG()
    assert run_digest(db, tg) is True and len(tg.sent) == 2
    assert run_digest(db, tg) is False and len(tg.sent) == 2
    assert run_digest(db, tg, force=True) is True and len(tg.sent) == 4
    assert "BẢN TIN SỰ KIỆN MIỀN NAM" in tg.sent[0][1]


def test_source_failure_alert(monkeypatch):
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "100")
    monkeypatch.setenv("ADMIN_CHAT_IDS", "1")
    db, tg = DB(":memory:"), FakeTG()
    for _ in range(2):
        db.health_fail("s1", "Nguồn 1", "Timeout")
    notify_source_failures(db, tg)
    assert tg.sent == []
    db.health_fail("s1", "Nguồn 1", "Timeout")
    notify_source_failures(db, tg)
    assert len(tg.sent) == 1 and tg.sent[0][0] == "1" and "Nguồn 1" in tg.sent[0][1]
    notify_source_failures(db, tg)  # không báo lặp
    assert len(tg.sent) == 1
    assert db.health_ok("s1", "Nguồn 1", 5) is True  # hồi phục


def test_tick_schedule(monkeypatch):
    from bot import runner
    monkeypatch.setenv("TELEGRAM_CHAT_IDS", "100")
    calls = []
    monkeypatch.setattr(runner, "run_collect", lambda *a, **k: (calls.append("collect"), a[0].set("last_collect", config.now().isoformat()))[0])
    fake_now = lambda h: datetime(2026, 10, 12, h, 5, tzinfo=config.TZ)  # Thứ Hai
    db, tg = seeded_db(), FakeTG()
    monkeypatch.setattr(config, "now", lambda: fake_now(6))
    runner.tick(db, tg)
    assert [t for _, t in tg.sent if "BẢN TIN" in t] == []         # 06:05 chưa đến giờ
    monkeypatch.setattr(config, "now", lambda: fake_now(7))
    runner.tick(db, tg)
    assert len([t for _, t in tg.sent if "BẢN TIN" in t]) == 1       # 07:05 gửi
    assert len(tg.docs) == 1                                          # thứ Hai kèm Excel tuần
    runner.tick(db, tg)
    assert len([t for _, t in tg.sent if "BẢN TIN" in t]) == 1       # không gửi lặp


def test_sources_config():
    s = load_sources()
    ids = [x["id"] for x in s]
    assert len(ids) == len(set(ids)) and len(s) >= 5
    assert sum(1 for x in s if x.get("is_gnews")) == 8
    assert all(not x["id"].startswith("mau-") for x in s)  # nguồn mẫu mặc định tắt


def test_failure_alert_is_batched(monkeypatch):
    monkeypatch.setenv("ADMIN_CHAT_IDS", "1")
    db, tg = DB(":memory:"), FakeTG()
    for i in range(10):
        for _ in range(3):
            db.health_fail(f"s{i}", f"Nguồn {i}", "HTTPError: 403")
    notify_source_failures(db, tg)
    assert len(tg.sent) == 1 and "10 nguồn" in tg.sent[0][1]
